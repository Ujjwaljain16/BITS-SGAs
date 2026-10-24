"""
Generate FoodExpress Assignment Document (DOCX & PDF) and Final Zip Package
Student: UJJWAL JAIN
ID: 2024EB02528
Subject: Software Development Practices
"""

import os
import zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Image as RLImage,
    PageBreak, KeepTogether, HRFlowable, Table as RLTable, TableStyle as RLTableStyle
)
from PIL import Image

DIR = os.path.dirname(os.path.abspath(__file__))


def read_puml(filename):
    with open(os.path.join(DIR, filename), "r", encoding="utf-8") as f:
        return f.read().strip()


def build_word_document(docx_path):
    doc = Document()
    
    # Page setup - 0.75 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Header / Meta
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("BITS Pilani - Work Integrated Learning Programmes\n")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(16)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    run_sub = p_title.add_run("Software Development Practices\nStaff Graded Assignment 1\n")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(12)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(30, 58, 138)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(18)
    run_meta = p_meta.add_run("Student Name: UJJWAL JAIN | Student ID: 2024EB02528\nSystem Design: FoodExpress Online Food Ordering and Delivery Platform")
    run_meta.font.name = "Arial"
    run_meta.font.size = Pt(10.5)
    run_meta.font.color.rgb = RGBColor(51, 65, 85)

    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 58, 138)
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(3)
        r = h.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = RGBColor(15, 118, 110)
        return h

    def add_body(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(30, 41, 59)
        return p

    def add_code(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(text)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(3, 105, 161)
        return p

    # --- Q1 ---
    add_h1("1. Functional and Non-Functional Requirements (4 Marks)")
    add_h2("Functional Requirements (FR)")
    add_body("FR1 (Restaurant and Menu Browsing): The system shall enable customers to search, filter by cuisine/location, and view restaurant menus with real-time pricing and availability status.")
    add_body("FR2 (Cart Management and Order Checkout): The system shall allow customers to add items to cart, customize quantities, add special preparation instructions, and place orders.")
    add_body("FR3 (Payment Processing): The system shall integrate with external payment gateways to process electronic transactions (Cards, UPI, Net Banking) and issue digital receipts.")
    add_body("FR4 (Restaurant Order Management): The system shall deliver orders in real time to the selected restaurant, allowing kitchen staff to accept or reject orders and update preparation status.")
    add_body("FR5 (Delivery Assignment and Live Tracking): The system shall automatically assign confirmed orders to nearby delivery partners and provide customers with live delivery status tracking.")

    add_h2("Non-Functional Requirements (NFR)")
    add_body("NFR1 (Performance and Latency): Menu search queries and checkout processing shall return responses within 1.5 seconds at the 95th percentile under standard operating load.")
    add_body("NFR2 (Availability and Uptime): The core customer ordering platform shall maintain 99.9% uptime, particularly during peak meal rush periods (12:00 to 14:30 and 19:00 to 22:30).")
    add_body("NFR3 (Security and Data Protection): Payment transactions must comply with PCI-DSS Level 1 standards. Customer personal data, passwords, and tokens must be encrypted using AES-256 and bcrypt, with TLS 1.3 enforced on all endpoints.")
    add_body("NFR4 (Scalability and Concurrency): The backend architecture shall scale horizontally to support at least 10,000 concurrent shopping carts and active checkout sessions without transaction deadlocks.")

    # --- Q2 ---
    add_h1("2. Use Case: Place a Food Order (3 Marks)")
    uc_text = (
        "Use Case Name: Place a Food Order\n"
        "Scope: FoodExpress Online Ordering System\n"
        "Level: User Goal\n"
        "Primary Actor: Customer\n"
        "Secondary Actors: Payment Gateway, Restaurant, Delivery Partner, Notification Service\n\n"
        "Stakeholders and Interests:\n"
        "- Customer: Wants a quick, secure checkout, accurate order itemization, and clear delivery ETA.\n"
        "- Restaurant: Wants verified payments before starting food preparation and clear special instructions.\n"
        "- Delivery Partner: Wants prompt notification of orders ready for pickup.\n"
        "- FoodExpress Platform: Wants verified transactions, reliable order fulfillment, and low dispute rates.\n\n"
        "Preconditions:\n"
        "1. Customer is logged in to an active account with a verified delivery address.\n"
        "2. Cart contains at least one food item from an active, open restaurant.\n\n"
        "Success Postconditions:\n"
        "1. Order record is created and stored in Confirmed state.\n"
        "2. Payment is successfully authorized and captured.\n"
        "3. Order details are transmitted to the restaurant kitchen.\n"
        "4. Order confirmation notification with ETA is sent to Customer.\n\n"
        "Failure Postconditions:\n"
        "1. No financial debit occurs, or full refund is initiated immediately.\n"
        "2. Customer cart items remain preserved, and clear error reason is shown.\n\n"
        "Main Success Scenario (Flow of Events):\n"
        "1. Customer navigates to checkout screen and reviews cart items, delivery address, and pricing.\n"
        "2. Customer selects payment method and clicks 'Place Order'.\n"
        "3. System generates an order with status 'Payment Pending' and redirects to Payment Gateway.\n"
        "4. Customer enters payment credentials and confirms transaction.\n"
        "5. Payment Gateway authorizes transaction and returns payment success token to System.\n"
        "6. System updates order payment status to 'Payment Received' and forwards order to Restaurant.\n"
        "7. Restaurant reviews order and clicks 'Accept Order'.\n"
        "8. System updates order status to 'Confirmed' and sets estimated preparation time.\n"
        "9. System dispatches confirmation notification with Order ID and ETA to Customer.\n"
        "10. System registers order in the delivery dispatch pool for partner assignment.\n\n"
        "Extensions / Alternative Flows:\n"
        "3a. Item Out of Stock: System alerts Customer, recalculates total, and prompts to modify cart.\n"
        "5a. Payment Failed: Payment Gateway declines; System displays error message and prompts retry.\n"
        "7a. Restaurant Rejects Order: Restaurant indicates capacity full or item unavailable; System cancels order, triggers full refund, and notifies Customer."
    )
    add_body(uc_text)

    # --- Q3 ---
    add_h1("3. Sequence Diagram: Placing a Food Order (3 Marks)")
    add_body("The sequence diagram illustrates interactions across Customer, Mobile App UI, Order Controller, Payment Service, Restaurant Service, and Notification Service.")
    img_seq = os.path.join(DIR, "sequence_diagram.png")
    if os.path.exists(img_seq):
        doc.add_picture(img_seq, width=Inches(6.8))
    add_h2("PlantUML Source Code (Sequence Diagram)")
    add_code(read_puml("sequence_diagram.puml"))

    # --- Q4 ---
    add_h1("4. Activity Diagram: Complete Food Ordering Workflow (2 Marks)")
    add_body("The activity diagram models the end-to-end food ordering workflow across Customer, FoodExpress System, Restaurant, and Delivery Partner, highlighting key decision points.")
    img_act = os.path.join(DIR, "activity_diagram.png")
    if os.path.exists(img_act):
        doc.add_picture(img_act, width=Inches(6.8))
    add_h2("PlantUML Source Code (Activity Diagram)")
    add_code(read_puml("activity_diagram.puml"))

    # --- Q5 ---
    add_h1("5. State Diagram: Order Lifecycle (3 Marks)")
    add_body("The state machine diagram illustrates all valid states and transitions for an Order object from creation to delivery or cancellation.")
    img_state = os.path.join(DIR, "state_diagram.png")
    if os.path.exists(img_state):
        doc.add_picture(img_state, width=Inches(6.0))
    add_h2("PlantUML Source Code (State Diagram)")
    add_code(read_puml("state_diagram.puml"))

    # --- Q6 ---
    add_h1("6. Class Diagram: FoodExpress System (5 Marks)")
    add_body("The class diagram defines the structural domain model including Customer, Restaurant, Menu, FoodItem, Order, OrderItem, Payment, and DeliveryPartner with explicit access modifiers, attributes, methods, and relationships.")
    img_cls = os.path.join(DIR, "class_diagram.png")
    if os.path.exists(img_cls):
        doc.add_picture(img_cls, width=Inches(6.8))
    add_h2("PlantUML Source Code (Class Diagram)")
    add_code(read_puml("class_diagram.puml"))

    doc.save(docx_path)
    print(f"Saved DOCX: {docx_path}")


def get_scaled_rl_image(filename, max_w=540, max_h=380):
    filepath = os.path.join(DIR, filename)
    with Image.open(filepath) as img:
        orig_w, orig_h = img.size
    aspect = orig_w / float(orig_h)
    w = min(max_w, orig_w)
    h = w / aspect
    if h > max_h:
        h = max_h
        w = h * aspect
    return RLImage(filepath, width=w, height=h)


def build_pdf_document(pdf_path):
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()

    header_title = ParagraphStyle(
        "HeaderTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=18,
        textColor=colors.HexColor("#0f172a"),
        alignment=1,
        spaceAfter=3
    )

    header_meta = ParagraphStyle(
        "HeaderMeta",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#1e3a8a"),
        alignment=1,
        spaceAfter=2
    )

    section_heading = ParagraphStyle(
        "SectionHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=15,
        textColor=colors.HexColor("#1e3a8a"),
        spaceBefore=10,
        spaceAfter=4
    )

    sub_heading = ParagraphStyle(
        "SubHeading",
        parent=styles["Heading3"],
        fontName="Helvetica-Bold",
        fontSize=9.5,
        leading=12,
        textColor=colors.HexColor("#0f766e"),
        spaceBefore=6,
        spaceAfter=2
    )

    body_text = ParagraphStyle(
        "BodyText",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=4
    )

    code_text = ParagraphStyle(
        "CodeText",
        parent=styles["Normal"],
        fontName="Courier",
        fontSize=7,
        leading=9,
        textColor=colors.HexColor("#0369a1"),
        backColor=colors.HexColor("#f8fafc"),
        borderPadding=4,
        spaceAfter=6
    )

    story = []

    # Header
    story.append(Paragraph("BITS Pilani - Work Integrated Learning Programmes", header_title))
    story.append(Paragraph("Software Development Practices | Staff Graded Assignment 1", header_meta))
    story.append(Paragraph("Student Name: UJJWAL JAIN | Student ID: 2024EB02528", header_meta))
    story.append(Paragraph("System Design Project: FoodExpress Online Food Ordering System", header_meta))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#1e3a8a"), spaceAfter=8))

    # Q1
    story.append(Paragraph("1. Functional and Non-Functional Requirements (4 Marks)", section_heading))
    story.append(Paragraph("<b>Functional Requirements (FR):</b>", sub_heading))
    story.append(Paragraph("• <b>FR1 (Restaurant and Menu Browsing):</b> The system shall enable customers to search, filter by cuisine/location, and view restaurant menus with real-time pricing and availability.", body_text))
    story.append(Paragraph("• <b>FR2 (Cart Management and Checkout):</b> The system shall allow customers to add items to cart, modify quantities, add special instructions, and submit orders for checkout.", body_text))
    story.append(Paragraph("• <b>FR3 (Payment Processing):</b> The system shall integrate with external payment gateways to process electronic transactions (Cards, UPI, Net Banking) and issue digital receipts.", body_text))
    story.append(Paragraph("• <b>FR4 (Restaurant Order Management):</b> The system shall notify restaurants in real time, allowing kitchen staff to accept or reject orders and update preparation status.", body_text))
    story.append(Paragraph("• <b>FR5 (Delivery Assignment and Live Tracking):</b> The system shall assign confirmed orders to nearby delivery partners and provide customers with live order tracking.", body_text))

    story.append(Paragraph("<b>Non-Functional Requirements (NFR):</b>", sub_heading))
    story.append(Paragraph("• <b>NFR1 (Performance and Latency):</b> Menu search and checkout processing shall return responses within 1.5 seconds at the 95th percentile under normal operating load.", body_text))
    story.append(Paragraph("• <b>NFR2 (Availability and Uptime):</b> The core ordering platform shall maintain 99.9% uptime, especially during peak meal rush hours (12:00 to 14:30 and 19:00 to 22:30).", body_text))
    story.append(Paragraph("• <b>NFR3 (Security and Data Protection):</b> Payment transactions must comply with PCI-DSS Level 1 standards. Customer personal data, passwords, and tokens must be encrypted using AES-256 and bcrypt, with TLS 1.3 enforced.", body_text))
    story.append(Paragraph("• <b>NFR4 (Scalability and Concurrency):</b> The backend architecture shall scale horizontally to support at least 10,000 concurrent shopping carts and active checkout sessions without transaction deadlocks.", body_text))

    story.append(PageBreak())

    # Q2
    story.append(Paragraph("2. Use Case: Place a Food Order (3 Marks)", section_heading))
    uc_lines = [
        "<b>Use Case Name:</b> Place a Food Order",
        "<b>Scope:</b> FoodExpress Online Ordering System | <b>Level:</b> User Goal",
        "<b>Primary Actor:</b> Customer",
        "<b>Secondary Actors:</b> Payment Gateway, Restaurant, Delivery Partner, Notification Service",
        "<b>Stakeholders and Interests:</b> Customer wants fast checkout and accurate delivery ETA. Restaurant wants verified payments before cooking. Delivery Partner wants prompt notification of orders ready for pickup. FoodExpress Platform wants verified transactions and low dispute rates.",
        "<b>Preconditions:</b> (1) Customer is logged in with a verified delivery address. (2) Cart contains at least one item from an active, open restaurant.",
        "<b>Success Postconditions:</b> Order record is persisted in Confirmed state. Payment is authorized. Kitchen ticket is dispatched to Restaurant. Order confirmation with ETA is sent to Customer.",
        "<b>Failure Postconditions:</b> No financial debit occurs (or immediate refund is issued). Customer cart items remain preserved, and clear error reason is shown.",
        "<b>Main Success Scenario (Nominal Flow):</b><br/>"
        "1. Customer navigates to checkout screen and reviews cart items, delivery address, and pricing.<br/>"
        "2. Customer selects payment method and clicks 'Place Order'.<br/>"
        "3. System generates an order with status 'Payment Pending' and redirects to Payment Gateway.<br/>"
        "4. Customer enters payment credentials and confirms transaction.<br/>"
        "5. Payment Gateway authorizes transaction and returns payment success token to System.<br/>"
        "6. System updates order payment status to 'Payment Received' and forwards order to Restaurant.<br/>"
        "7. Restaurant reviews order and clicks 'Accept Order'.<br/>"
        "8. System updates order status to 'Confirmed' and sets estimated preparation time.<br/>"
        "9. System dispatches confirmation notification with Order ID and ETA to Customer.<br/>"
        "10. System registers order in delivery dispatch pool for partner assignment.",
        "<b>Extensions / Alternative Flows:</b><br/>"
        "3a. Item Out of Stock: System alerts Customer, recalculates total, and prompts to modify cart.<br/>"
        "5a. Payment Failed: Payment Gateway declines; System displays error message and prompts retry.<br/>"
        "7a. Restaurant Rejects Order: Restaurant indicates capacity full; System cancels order, triggers full refund, and notifies Customer."
    ]
    for line in uc_lines:
        story.append(Paragraph(line, body_text))

    story.append(PageBreak())

    # Q3
    story.append(Paragraph("3. Sequence Diagram: Placing a Food Order (3 Marks)", section_heading))
    story.append(Paragraph("The sequence diagram illustrates interactions across Customer, Mobile App UI, Order Controller, Payment Service, Restaurant Service, and Notification Service.", body_text))
    story.append(get_scaled_rl_image("sequence_diagram.png", max_w=530, max_h=360))
    story.append(Spacer(1, 4))
    story.append(Paragraph("<b>PlantUML Source Code:</b>", sub_heading))
    story.append(Paragraph(read_puml("sequence_diagram.puml").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>"), code_text))

    story.append(PageBreak())

    # Q4
    story.append(Paragraph("4. Activity Diagram: Complete Food Ordering Workflow (2 Marks)", section_heading))
    story.append(Paragraph("The activity diagram models the end-to-end food ordering workflow across Customer, FoodExpress System, Restaurant, and Delivery Partner, highlighting key decision points.", body_text))
    story.append(get_scaled_rl_image("activity_diagram.png", max_w=530, max_h=360))
    story.append(Spacer(1, 4))
    story.append(Paragraph("<b>PlantUML Source Code:</b>", sub_heading))
    story.append(Paragraph(read_puml("activity_diagram.puml").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>"), code_text))

    story.append(PageBreak())

    # Q5
    story.append(Paragraph("5. State Diagram: Order Lifecycle (3 Marks)", section_heading))
    story.append(Paragraph("The state machine diagram illustrates all valid states and transitions for an Order object from creation to delivery or cancellation.", body_text))
    story.append(get_scaled_rl_image("state_diagram.png", max_w=500, max_h=300))
    story.append(Spacer(1, 4))
    story.append(Paragraph("<b>PlantUML Source Code:</b>", sub_heading))
    story.append(Paragraph(read_puml("state_diagram.puml").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>"), code_text))

    story.append(PageBreak())

    # Q6
    story.append(Paragraph("6. Class Diagram: FoodExpress System (5 Marks)", section_heading))
    story.append(Paragraph("The class diagram defines the structural domain model including Customer, Restaurant, Menu, FoodItem, Order, OrderItem, Payment, and DeliveryPartner with explicit access modifiers, attributes, methods, and relationships.", body_text))
    story.append(get_scaled_rl_image("class_diagram.png", max_w=530, max_h=360))
    story.append(Spacer(1, 4))
    story.append(Paragraph("<b>PlantUML Source Code:</b>", sub_heading))
    story.append(Paragraph(read_puml("class_diagram.puml").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>"), code_text))

    doc.build(story)
    print(f"Saved PDF: {pdf_path}")


def package_submission(zip_name):
    zip_path = os.path.join(DIR, zip_name)
    files_to_pack = [
        "sequence_diagram.puml",
        "sequence_diagram.png",
        "activity_diagram.puml",
        "activity_diagram.png",
        "state_diagram.puml",
        "state_diagram.png",
        "class_diagram.puml",
        "class_diagram.png",
        "Ujjwal_Jain_2024EB02528_SG1.docx",
        "Ujjwal_Jain_2024EB02528_SG1.pdf"
    ]
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        for f in files_to_pack:
            p = os.path.join(DIR, f)
            if os.path.exists(p):
                z.write(p, arcname=f)
                print(f"Added to zip: {f}")
            else:
                print(f"[WARNING] Missing: {f}")
    print(f"\nFinal Zip Package Created: {zip_path}")


if __name__ == "__main__":
    docx_file = os.path.join(DIR, "Ujjwal_Jain_2024EB02528_SG1.docx")
    pdf_file = os.path.join(DIR, "Ujjwal_Jain_2024EB02528_SG1.pdf")
    
    print("Building Word Document (.docx)...")
    build_word_document(docx_file)
    
    print("Building PDF Document (.pdf)...")
    build_pdf_document(pdf_file)
    
    print("Creating Submission Zip Archive...")
    package_submission("Ujjwal_Jain_2024EB02528_SG1.zip")
